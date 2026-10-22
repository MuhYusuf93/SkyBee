"""Independent readers for public M5.1 CSV/OpenXML artifacts, never store data."""
import csv
import json
import sys
from pathlib import Path
from zipfile import ZipFile
import openpyxl

root = Path(sys.argv[1] if len(sys.argv) > 1 else 'build/m51-report-reader')
with (root / 'report.csv').open(encoding='utf-8-sig', newline='') as source:
    csv_rows = list(csv.DictReader(source))
book = openpyxl.load_workbook(root / 'report.xlsx', read_only=True, data_only=False)
assert book.sheetnames == ['Summary', 'Sales', 'Payments', 'Items', 'Current inventory']
xlsx_rows = []
for sheet in book:
    iterator = sheet.iter_rows()
    header = [cell.value for cell in next(iterator)]
    for cells in iterator:
        assert all(cell.data_type != 'f' for cell in cells), 'formula cell'
        xlsx_rows.append(dict(zip(header, (cell.value for cell in cells))))
book.close()
for rows in [csv_rows, xlsx_rows]:
    assert len(rows) == 129
    summary = {r['metric']: r['value'] for r in rows if r['section'] == 'summary'}
    assert summary['net_sales'] == '12000'
    assert summary['net_payments'] == '11000'
    assert summary['cost'] == '9000000000'
    assert summary['current_receivables'] == '1000'
    sales = [r for r in rows if r['section'] == 'sales']
    payments = [r for r in rows if r['section'] == 'payments']
    integer = lambda v: int(str(v).removeprefix("'"))
    assert sum(integer(r['amount_minor']) for r in sales) == 12000
    assert sum(integer(r['cost_micros']) for r in sales) == 9000000000
    assert sum(integer(r['amount_minor']) for r in payments) == 11000
    assert len([r for r in rows if r['section'] == 'current_inventory']) == 106
    goods = next(r for r in rows if r['section'] == 'current_inventory' and r['barcode'] == '000123')
    assert '"Barang"\nEnglish' in goods['name']
    assert str(goods['stock_pcs']) == '9'
    assert int(goods['pcs_per_pack']) == 40
assert next(r for r in csv_rows if r['barcode'] == '000123')['name'].startswith("'=SUM")
assert next(r for r in xlsx_rows if r['barcode'] == '000123')['name'].startswith('=SUM')
assert all(isinstance(r['amount_minor'], str) for r in xlsx_rows if r['section'] in ('sales', 'payments', 'items'))
with ZipFile(root / 'report.xlsx') as archive:
    assert archive.testzip() is None
precision = openpyxl.load_workbook(root / 'precision/report.xlsx', read_only=True, data_only=False)
sheet = precision['Exact']
values = list(sheet.iter_rows())
row = dict(zip([cell.value for cell in values[0]], values[1]))
assert row['amount_minor'].value == '9223372036854775807'
assert row['amount_decimal'].value == '92233720368547758.07'
assert row['quantity'].value == '9223372036854775807'
assert row['name'].value == '=1+2' and row['name'].data_type != 'f'
precision.close()
print(json.dumps({'csv': 'PASS', 'xlsx_openpyxl': 'PASS', 'rows': 129,
                  'full_period_totals': 'PASS', 'barcode_zeroes': 'PASS',
                  'formula_and_multiline': 'PASS', 'typed_exact_money': 'PASS'}))
